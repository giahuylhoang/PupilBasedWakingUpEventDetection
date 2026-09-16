"""Tests for vasomotion_app.app using Streamlit's AppTest harness.

The five `test_app_*` functions below are the widget-scaffold smoke tests
from task-8-brief.md, kept verbatim: they cover the "no file uploaded yet"
state, which is all AppTest can drive for a real `st.file_uploader` in some
Streamlit versions -- but streamlit 1.63.0 (pinned in requirements.txt) does
support `FileUploader.set_value(...)`, confirmed directly against this
version's `streamlit.testing.v1.element_tree` source before relying on it
(rather than assumed from the brief's more cautious note). Every test past
those five therefore drives real file uploads and asserts on behaviour, not
just the absence of an exception:

- uploading the real sample workbook and producing full results
- the frequency-axis radio actually changing a rendered figure's scale
  (Defect 1's regression guard)
- the band-edge warning appearing, informatively, for the default settings
  (Defect 3)
- a NaN band value being explained rather than left as a silent blank
- the live sampling-rate readout showing all four numbers the spec requires
- a settings-JSON round trip genuinely restoring modified widget values
  (not merely displaying the file back at the user), including the one
  widget (the band-edge data_editor) that cannot be restored via
  st.session_state and is instead handled with a rotating widget key
- two crash-shaped inputs (a malformed settings file, an out-of-range NW)
  degrading to an error message or a clamp instead of an uncaught exception

Fix round 1 (review findings) added four more:

- selecting a preset after a manual edit updates the Advanced widgets'
  *displayed* values, not just the numbers computed underneath them
  (Finding 1)
- no module except app.py imports streamlit, discovered dynamically so a
  module added later is covered automatically (Finding 2)
- periodogram/multitaper_thomson (no `nperseg`) still get a band-edge
  warning when one is due, via the trace-length fallback (Finding 3)
- the NaN-band explanation does not appear when nothing is actually NaN

Fix round 2 (a defect introduced by fix round 1's own preset refactor)
added two more: `window`/`nw`/`k` used to fall back to a hardcoded literal
whenever their widget wasn't instantiated for the current method (e.g.
"Raw periodogram" sets window="boxcar", but periodogram never shows the
Window function widget, so parameters.json reported "hann"). Covered by a
test selecting that preset and a second test covering the latent form
(non-default nw/k surviving export under a method that hides their
widget too).

`st.pyplot` renders a static image with no queryable scale, and
`streamlit.testing.v1` has no typed accessor for `st.data_editor`, so
app.py stashes the actual Figure objects, the resolved band dict, and the
resolved PSDParams into `st.session_state["_test_figures"]` /
`["_test_band_defs"]` / `["_test_psd_params"]` purely so these tests can
inspect real objects rather than rendered pixels or a downloaded file's
bytes (which `st.download_button` does not expose to AppTest either).
"""
from __future__ import annotations

import inspect
import os

import openpyxl
import pytest
from streamlit.testing.v1 import AppTest

from vasomotion_app import app as app_module
from vasomotion_app import bundle
from vasomotion_app import preprocess
from vasomotion_app.psd_methods import PSDParams

APP_PATH = os.path.join(os.path.dirname(__file__), "..", "app.py")
SAMPLE_PATH = os.path.join(
    os.path.dirname(__file__), "..", "..", "data", "Tim",
    "Vasomotion analysis compilation_Bal K.xlsx",
)
# The workbook has no sampling rate of its own -- 2.0 Hz is this suite's own
# fixed choice, matching test_end_to_end.py, not a fact read from the file.
FS = 2.0

requires_sample_file = pytest.mark.skipif(
    not os.path.exists(SAMPLE_PATH), reason="sample workbook not present in this checkout",
)


# --- small widget-lookup helpers --------------------------------------
# AppTest exposes widgets as flat, type-grouped lists (at.sidebar.number_input,
# at.sidebar.selectbox, ...) regardless of nesting inside st.expander, so
# every widget here is found by its exact label rather than by position.

def _one(matches, description):
    assert len(matches) == 1, f"expected exactly one {description}, found {len(matches)}"
    return matches[0]


def _number_input(at, label):
    return _one([ni for ni in at.sidebar.number_input if ni.label == label], f"number_input {label!r}")


def _selectbox(at, label):
    return _one([sb for sb in at.sidebar.selectbox if sb.label == label], f"selectbox {label!r}")


def _radio(at, options):
    return _one(
        [r for r in at.sidebar.radio if list(r.options) == options],
        f"radio with options {options!r}",
    )


def _checkbox(at, label):
    return _one([c for c in at.sidebar.checkbox if c.label == label], f"checkbox {label!r}")


def _excel_uploader(at):
    return _one(
        [fu for fu in at.sidebar.file_uploader if "Excel" in fu.label], "the Excel file_uploader"
    )


def _settings_uploader(at):
    return _one(
        [fu for fu in at.sidebar.file_uploader if "parameters.json" in fu.label],
        "the settings file_uploader",
    )


def _fresh_app() -> AppTest:
    at = AppTest.from_file(APP_PATH, default_timeout=60)
    at.run()
    return at


def _upload_sample_and_configure(at: AppTest) -> AppTest:
    """Upload the real sample workbook and fill in the two required fields
    (fs=2.0 Hz, signal is Area), in one rerun, matching how a user staging
    several changes before Streamlit next reruns the script would land."""
    with open(SAMPLE_PATH, "rb") as f:
        content = f.read()
    _excel_uploader(at).set_value((
        "Vasomotion analysis compilation_Bal K.xlsx", content,
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ))
    _number_input(at, "Sampling rate (Hz)").set_value(FS)
    _radio(at, ["Area", "Diameter"]).set_value("Area")
    at.run()
    return at


# --- widget-scaffold smoke tests (task-8-brief.md, verbatim) ------------

def test_app_runs_without_exception():
    at = AppTest.from_file(APP_PATH)
    at.run()
    assert not at.exception


def test_app_title():
    at = AppTest.from_file(APP_PATH)
    at.run()
    assert at.title[0].value == "Vasomotion PSD Tool"


def test_app_shows_upload_prompt_when_no_file():
    at = AppTest.from_file(APP_PATH)
    at.run()
    info_texts = [i.value for i in at.info]
    assert any("Upload a vasomotion Excel file" in t for t in info_texts)


def test_app_signal_type_radio_has_no_default():
    at = AppTest.from_file(APP_PATH)
    at.run()
    radios = [r for r in at.sidebar.radio if list(r.options) == ["Area", "Diameter"]]
    assert len(radios) == 1
    assert radios[0].value is None


def test_app_preset_dropdown_has_three_presets_plus_custom():
    at = AppTest.from_file(APP_PATH)
    at.run()
    preset_boxes = [sb for sb in at.sidebar.selectbox if "Gordon lab default" in list(sb.options)]
    assert len(preset_boxes) == 1
    assert set(preset_boxes[0].options) == {
        "(custom)", "Gordon lab default", "Raw periodogram", "Long window",
    }


# --- behavioural coverage ------------------------------------------------

@requires_sample_file
def test_uploading_sample_workbook_produces_results():
    at = _upload_sample_and_configure(_fresh_app())
    assert not at.exception
    assert "Group-average spectra" in [h.value for h in at.header]
    assert "Band summary" in [h.value for h in at.header]
    assert "Per-animal spectra" in [h.value for h in at.header]

    figures = at.session_state["_test_figures"]
    assert set(figures["group_average"].keys()) == {"SmFP", "SpVB"}
    assert set(figures["per_animal"].keys()) == {"SmFP", "SpVB"}

    band_summary_tables = [
        d.value for d in at.dataframe
        if list(d.value.columns) == ["animal", "sex", "region", "condition", "ultra_low", "myogenic"]
    ]
    assert len(band_summary_tables) == 1
    assert len(band_summary_tables[0]) == 24  # 4 sheets x 6 animals -- test_end_to_end.py

    download_labels = {b.label for b in at.download_button}
    assert {
        "Download band summary (CSV)", "Download full PSD (CSV)",
        "Download per-trace statistics (CSV)",
        "Download parameters.json", "Download everything (.zip)",
    } <= download_labels

    # The exported per-trace statistics must actually reach the screen with
    # both SD columns -- the whole point of the table is that a user can
    # pull variability out of the app, so a silently missing column would
    # look like a working export right up until someone needed it.
    trace_stats_tables = [
        d.value for d in at.dataframe
        if list(d.value.columns) == preprocess.TRACE_STATS_COLUMNS
    ]
    assert len(trace_stats_tables) == 1
    stats = trace_stats_tables[0]
    assert len(stats) == 24  # same 24 traces as the band summary above
    assert (stats["sd_normalized"] >= stats["sd_analyzed"]).all()
    assert (stats["sd_normalized"] > 0).all()
    # signal_type is "area" in this run, so mean_raw is in diameter units:
    # the workbook's raw values span 249..4856, whose square roots span
    # ~15.8..69.7. A mean_raw above that ceiling would mean the area ->
    # diameter conversion was skipped for this table.
    assert stats["mean_raw"].max() < 70.0


@requires_sample_file
def test_changing_frequency_axis_radio_changes_rendered_figure_scale():
    """Regression guard for Defect 1. figures.plot_group_average defaults to
    xscale="log", and bundle.run_pipeline has no xscale parameter (its
    return contract is locked to exactly {full_psd, band_summary, figures}
    by Task 7's end-to-end test) -- so without app.py re-rendering after
    run_pipeline, this radio changes parameters.json but no figure ever
    moves off "log", however it is set."""
    at = _upload_sample_and_configure(_fresh_app())

    figures_before = at.session_state["_test_figures"]["group_average"]
    assert len(figures_before) == 2
    assert all(fig.axes[0].get_xscale() == "log" for fig in figures_before.values())

    _radio(at, ["log", "linear"]).set_value("linear")
    at.run()
    assert not at.exception

    figures_after = at.session_state["_test_figures"]["group_average"]
    assert len(figures_after) == 2
    assert all(fig.axes[0].get_xscale() == "linear" for fig in figures_after.values())
    # No control exposes yscale; the fix must not touch it.
    assert all(fig.axes[0].get_yscale() == "log" for fig in figures_after.values())


@requires_sample_file
def test_band_edge_warning_appears_with_default_preset():
    """Defect 3. At fs=2.0 with nperseg=256 the grid spacing is exactly
    0.0078125 Hz, so the Gordon lab default preset's own ultra_low band
    (lower edge 0.005 Hz) sits below the resolvable minimum -- confirmed by
    Task 7 as correct behaviour, not a bug, and by test_bands.py's own
    test_validate_band_edges_warns_below_resolution with these exact
    numbers. It must fire and must read as informative, not alarming."""
    at = _upload_sample_and_configure(_fresh_app())
    _selectbox(at, "Preset").set_value("Gordon lab default")
    at.run()
    assert not at.exception

    warnings = [w.value for w in at.warning]
    assert any("ultra_low" in w and "resolvable minimum" in w for w in warnings)

    info_texts = [i.value for i in at.info]
    assert any("still computed" in t for t in info_texts)


@requires_sample_file
def test_nan_band_value_is_explained_in_ui():
    at = _upload_sample_and_configure(_fresh_app())
    # Push Max frequency below the myogenic band (0.04-0.15 Hz) entirely, so
    # fewer than two frequency points fall inside it and integrate_bands
    # (bands.py) returns NaN for that column.
    _number_input(at, "Max frequency (Hz)").set_value(0.03)
    at.run()
    assert not at.exception

    band_summary_tables = [
        d.value for d in at.dataframe if "myogenic" in getattr(d.value, "columns", [])
    ]
    assert len(band_summary_tables) == 1
    assert band_summary_tables[0]["myogenic"].isna().all()

    info_texts = [i.value for i in at.info]
    assert any("'myogenic'" in t and "blank" in t for t in info_texts)


@requires_sample_file
def test_live_readout_shows_sample_count_duration_and_window_count():
    at = _upload_sample_and_configure(_fresh_app())
    readout = [
        i.value for i in at.info
        if "samples" in i.value and "duration" in i.value
    ]
    assert len(readout) == 1
    text = readout[0]
    assert "1173" in text  # sample count, from the real workbook
    assert "Lowest resolvable frequency" in text
    assert "windows to average" in text


def test_settings_round_trip_restores_modified_values():
    at = _fresh_app()

    # Move several controls away from their defaults first, so restoring
    # can be shown to actually change something rather than coincide with
    # values already in place.
    _number_input(at, "Sampling rate (Hz)").set_value(9.0)
    _radio(at, ["Area", "Diameter"]).set_value("Diameter")
    at.run()
    assert _number_input(at, "Sampling rate (Hz)").value == 9.0
    assert _radio(at, ["Area", "Diameter"]).value == "Diameter"

    params = bundle.build_parameters_dict(
        signal_type="area", normalize="zscore", detrend_method="constant",
        highpass_cutoff=None, highpass_order=4, trim_seconds=3.0,
        psd_params=PSDParams(
            method="welch", nperseg=128, noverlap=64, window="hamming",
            nw=3.0, k=5, zero_pad=2, scaling="spectrum",
        ),
        max_freq=0.35, xscale="linear", bands={"lowband": (0.01, 0.2)}, fs=5.5,
    )
    payload = bundle.parameters_to_json(params).encode("utf-8")
    _settings_uploader(at).set_value(("mysettings.json", payload, "application/json"))
    at.run()
    assert not at.exception
    assert any("loaded and applied" in s.value for s in at.success)

    assert _number_input(at, "Sampling rate (Hz)").value == 5.5
    assert _radio(at, ["Area", "Diameter"]).value == "Area"
    assert _selectbox(at, "Preset").value == "(custom)"
    assert _selectbox(at, "Normalise").value == "zscore"
    assert _selectbox(at, "Detrend").value == "constant"
    assert _checkbox(at, "High-pass filter").value is False
    assert _number_input(at, "Trim first N seconds").value == 3.0
    assert _selectbox(at, "Method").value == "welch"
    assert _number_input(at, "Window length (samples)").value == 128
    assert _number_input(at, "Overlap (samples)").value == 64
    assert _selectbox(at, "Window function").value == "hamming"
    assert _selectbox(at, "Zero-padding").value == 2
    assert _selectbox(at, "Scaling").value == "spectrum"
    assert _number_input(at, "Max frequency (Hz)").value == 0.35
    assert _radio(at, ["log", "linear"]).value == "linear"
    # The band-edge table (st.data_editor) has no typed AppTest accessor and
    # refuses session_state writes outright, so app.py restores it via a
    # rotating widget key instead; this is the only way to confirm that
    # path actually worked rather than silently no-opping.
    assert at.session_state["_test_band_defs"] == {"lowband": (0.01, 0.2)}

    # A settings file stays "uploaded" across reruns until cleared; a later,
    # unrelated manual edit must not be silently overwritten by it again.
    _number_input(at, "Sampling rate (Hz)").set_value(42.0)
    at.run()
    assert not at.exception
    assert _number_input(at, "Sampling rate (Hz)").value == 42.0


def test_malformed_settings_file_shows_error_not_crash():
    at = _fresh_app()
    _settings_uploader(at).set_value(("bad.json", b"{not valid json", "application/json"))
    at.run()
    assert not at.exception
    assert any("Could not load settings file" in e.value for e in at.error)


def test_nw_minimum_prevents_invalid_k_range_crash():
    """K's range is 1..2*NW-1; any NW below 1.0 makes that range empty,
    which crashed the whole app with an uncaught StreamlitAPIException
    (reproduced while building app.py, since the brief's reference code
    allows NW down to 0.5). NW's minimum is now 1.0, the true lower bound
    for a non-empty K range, so this must degrade to a clamp, not a crash."""
    at = _fresh_app()
    _number_input(at, "NW").set_value(0.5)
    at.run()
    assert not at.exception
    assert _number_input(at, "NW").value >= 1.0


# --- fix round 1 (review findings) ---------------------------------------

def test_selecting_preset_updates_method_selectbox_after_manual_edit():
    """Regression guard for Finding 1. Reproduced by the reviewer: set
    Method to "welch" by hand, then select "Gordon lab default" -- before
    this fix the Method selectbox still read "welch" while the results
    demonstrably came from multitaper_welch (parameters.json was always
    correct; only the sidebar's displayed value was stale). Selecting a
    preset must make the widget redraw showing what will actually run."""
    at = _fresh_app()
    _selectbox(at, "Method").set_value("welch")
    at.run()
    assert _selectbox(at, "Method").value == "welch"

    _selectbox(at, "Preset").set_value("Gordon lab default")
    at.run()
    assert not at.exception
    assert _selectbox(at, "Method").value == "multitaper_welch"
    assert _number_input(at, "Window length (samples)").value == 256
    assert _number_input(at, "Overlap (samples)").value == 128
    assert _number_input(at, "NW").value == 3.0


def test_no_module_except_app_imports_streamlit():
    """Architectural guard for Finding 2. Design spec: "The analysis core
    does not import Streamlit." Discovers vasomotion_app/*.py dynamically
    (not a hardcoded list) so a module added later is covered
    automatically. Scoped to the top-level package files -- not
    vasomotion_app/tests/, which legitimately imports streamlit.testing to
    drive AppTest; that is test infrastructure, not the analysis core this
    constraint protects. Reproduced against the actual bug: adding `import
    streamlit` to a scratch copy of figures.py makes this fail, naming
    figures.py, while the rest of the suite stays green (see task-8-report.md,
    Fix round 1)."""
    import ast
    import pathlib

    package_dir = pathlib.Path(__file__).resolve().parent.parent
    assert package_dir.name == "vasomotion_app"

    def _imports_streamlit(path: pathlib.Path) -> bool:
        tree = ast.parse(path.read_text(), filename=str(path))
        for node in ast.walk(tree):
            if isinstance(node, ast.Import):
                if any(a.name == "streamlit" or a.name.startswith("streamlit.") for a in node.names):
                    return True
            elif isinstance(node, ast.ImportFrom):
                if node.module and (node.module == "streamlit" or node.module.startswith("streamlit.")):
                    return True
        return False

    offenders = [
        path.name for path in sorted(package_dir.glob("*.py"))
        if path.name != "app.py" and _imports_streamlit(path)
    ]
    assert offenders == [], f"streamlit imported outside app.py in: {offenders}"


@requires_sample_file
def test_band_edge_warning_fires_under_periodogram_for_out_of_range_band():
    """Regression guard for Finding 3. periodogram/multitaper_thomson have
    no `nperseg` (they use the whole trace as their one window), so
    app.py's window_length fallback (trace length instead of nperseg) must
    still be consulted for those methods, not silently skipped -- reverting
    to the brief's original `if psd_params.nperseg is not None:` gate
    around the whole validate_band_edges call would make this go silent.
    At fs=2.0 Hz with the real sample file's 1173 samples, the whole-trace
    resolvable minimum is 2/1173 ~= 0.0017 Hz; a band edge below that must
    still warn under periodogram."""
    at = _upload_sample_and_configure(_fresh_app())
    params = bundle.build_parameters_dict(
        signal_type="area", normalize="dx_x0_median", detrend_method="linear",
        highpass_cutoff=0.005, highpass_order=4, trim_seconds=0.0,
        psd_params=PSDParams(
            method="periodogram", nperseg=None, noverlap=None, window="boxcar",
            nw=3.0, k=5, zero_pad=1, scaling="density",
        ),
        max_freq=0.5, xscale="log", bands={"tiny": (0.0005, 0.001)}, fs=FS,
    )
    payload = bundle.parameters_to_json(params).encode("utf-8")
    _settings_uploader(at).set_value(("periodogram_settings.json", payload, "application/json"))
    at.run()
    assert not at.exception
    assert _selectbox(at, "Method").value == "periodogram"
    assert at.session_state["_test_band_defs"] == {"tiny": (0.0005, 0.001)}

    warnings = [w.value for w in at.warning]
    assert any("tiny" in w and "resolvable minimum" in w for w in warnings)


@requires_sample_file
def test_nan_explanation_does_not_appear_when_no_band_is_nan():
    """Mutation guard: the NaN-band explanation is gated on `if nan_bands:`
    (app.py) -- a mutation to `if True:` would leave the rest of the suite
    green, since every other NaN-related test only checks the explanation
    DOES appear when a band genuinely is NaN. Gordon lab default's own two
    bands are both fully resolvable on the real sample file (confirmed by
    test_end_to_end.py -- zero NaNs), so no explanation should appear."""
    at = _upload_sample_and_configure(_fresh_app())
    _selectbox(at, "Preset").set_value("Gordon lab default")
    at.run()
    assert not at.exception
    info_texts = [i.value for i in at.info]
    assert not any("blank cells" in t for t in info_texts)


# --- fix round 2 (a defect introduced by fix round 1's own refactor) ----

def test_raw_periodogram_preset_exports_true_window_not_hardcoded_fallback():
    """Regression guard for the Fix-round-2 defect, updated for fix-wave
    item 3 (see bundle.build_parameters_dict's docstring for the full
    rationale -- "this is the third round on this class"). The Window
    function widget only exists `if method == "welch"`; before round 2's
    fix, `window` fell back to a hardcoded "hann" literal whenever that
    widget wasn't instantiated, silently discarding whatever a preset or a
    loaded settings file actually declared -- "Raw periodogram" declares
    window="boxcar" (PRESETS["Raw periodogram"]["psd"]["window"]) while
    method="periodogram" never shows that widget. Round 2 fixed the
    WIDGET/PSDParams layer to keep remembering "boxcar" (asserted below,
    unchanged) -- but round 2 then also let that same "boxcar" leak into
    the EXPORTED parameters.json, which item 3 corrects: compute_psd never
    reads `window` for periodogram at all (it hardcodes a boxcar
    internally, taken from no field of `params`), so exporting any window
    value for this run -- "hann", "boxcar", or anything else -- would
    misreport a field that did not affect what ran. The exported value is
    now null, not a literal that happens to coincide with what silently
    ran."""
    at = _fresh_app()
    _selectbox(at, "Preset").set_value("Raw periodogram")
    at.run()
    assert not at.exception

    psd_params = at.session_state["_test_psd_params"]
    assert psd_params.method == "periodogram"
    assert psd_params.window == "boxcar"

    # The literal check the reviewer asked for: round-trip through the
    # actual parameters.json machinery, not just the PSDParams object.
    parameters = bundle.build_parameters_dict(
        signal_type="area", normalize="dx_x0_median", detrend_method="linear",
        highpass_cutoff=0.005, highpass_order=4, trim_seconds=0.0,
        psd_params=psd_params, max_freq=0.5, xscale="log",
        bands={"x": (0.01, 0.1)}, fs=FS,
    )
    reloaded = bundle.parameters_from_json(bundle.parameters_to_json(parameters))
    assert reloaded["psd"]["window"] is None


def test_legacy_settings_file_with_moot_nw_k_still_reloads_them():
    """The latent form of the same defect, updated for fix-wave item 3.
    The NW/K widgets only exist `if method in ("multitaper_welch",
    "multitaper_thomson")`; round 2 fixed the WIDGET layer to keep
    remembering a configured nw/k across a switch to a method that hides
    them. Item 3 now nulls nw/k at EXPORT time for a method that ignores
    them (here, "welch"), so the app itself can no longer produce a
    downloaded parameters.json with method="welch" and a non-null nw/k --
    which means this scenario can no longer be built via
    bundle.build_parameters_dict, unlike the original round-2 test. It can
    still arrive as a hand-edited or older-app-version settings file,
    though, and `_push_settings_into_widgets`'s reload logic (a distinct
    concern from what THIS app chooses to export -- see
    build_parameters_dict's docstring) makes no method-conditional
    judgement about whether to restore a field; it restores whatever is
    present and non-null. This constructs such a file directly (not via
    build_parameters_dict, which would null nw/k right back out) and
    confirms the widgets still pick up nw=7.0/k=9 rather than silently
    reporting the nw=3.0/k=5 widget defaults."""
    at = _fresh_app()
    params_in = {
        "app_version": bundle.APP_VERSION, "fs": FS, "signal_type": "area",
        "normalize": "dx_x0_median", "detrend": "linear", "highpass_cutoff": 0.005,
        "highpass_order": 4, "trim_seconds": 0.0,
        "psd": {
            "method": "welch", "nperseg": 256, "noverlap": 128, "window": "hamming",
            "nw": 7.0, "k": 9, "zero_pad": 1, "scaling": "density",
        },
        "max_freq": 0.5, "xscale": "log", "bands": {"x": [0.01, 0.1]},
    }
    payload = bundle.parameters_to_json(params_in).encode("utf-8")
    _settings_uploader(at).set_value(("nwk_settings.json", payload, "application/json"))
    at.run()
    assert not at.exception
    assert _selectbox(at, "Method").value == "welch"

    psd_params = at.session_state["_test_psd_params"]
    assert psd_params.nw == 7.0
    assert psd_params.k == 9


# --- final fix wave (whole-branch review) --------------------------------


@requires_sample_file
def test_spectrum_scaling_with_multitaper_shows_error_not_traceback():
    """CRITICAL fix-wave item 1's app-level wiring. multitaper_welch is the
    shipped default method, and 'spectrum' is a selectable Scaling option
    -- before the psd_methods.py fix, that combination silently computed
    power values ~1e26-1e31 too large, with no exception, no warning, no
    NaN. After compute_psd raises ValueError for it, this confirms app.py
    surfaces that as a plain st.error (proactively, before run_pipeline
    ever runs), not as an uncaught exception."""
    at = _upload_sample_and_configure(_fresh_app())
    assert _selectbox(at, "Method").value == "multitaper_welch"  # the shipped default
    _selectbox(at, "Scaling").set_value("spectrum")
    at.run()
    assert not at.exception
    errors = [e.value for e in at.error]
    assert any("scaling='spectrum' is undefined for method 'multitaper_welch'" in e for e in errors)


@requires_sample_file
def test_compare_methods_with_raw_periodogram_preset_does_not_crash():
    """CRITICAL fix-wave item 2. Exact three-click repro from the review:
    upload -> preset 'Raw periodogram' -> tick 'Compare methods'. The
    multiselect's shipped default is ["periodogram", "multitaper_welch"],
    and 'Raw periodogram' sets psd.nperseg=None (periodogram has no window
    length of its own) -- before this fix, building
    PSDParams(method="multitaper_welch", nperseg=None) and handing it to
    compute_psd hit a bare `int(None)` and raised an uncaught TypeError
    that the compare-methods loop's `except ValueError` could not catch,
    aborting the whole script before "Download parameters.json" and
    "Download everything (.zip)" ever rendered."""
    at = _upload_sample_and_configure(_fresh_app())
    _selectbox(at, "Preset").set_value("Raw periodogram")
    at.run()
    assert not at.exception

    _checkbox(at, "Compare methods").set_value(True)
    at.run()
    assert not at.exception

    assert "Compare methods" in [h.value for h in at.header]

    # Stronger than "did not crash": the comparison must actually succeed,
    # not merely fail gracefully. compute_psd's own None-guard (turning the
    # crash into a caught ValueError) is necessary but not sufficient on
    # its own -- without app.py also substituting a valid window length for
    # the compared Welch-family method, this would land here with a caught
    # "Could not compute 'multitaper_welch' for comparison" error and no
    # percent_diff table, which is a regression this assertion catches
    # that "no exception" alone would not.
    assert not any("Could not compute" in e.value for e in at.error)
    comparison_tables = [
        d.value for d in at.dataframe if "percent_diff" in getattr(d.value, "columns", [])
    ]
    assert len(comparison_tables) == 1

    download_labels = {b.label for b in at.download_button}
    assert {
        "Download band summary (CSV)", "Download full PSD (CSV)",
        "Download parameters.json", "Download everything (.zip)",
    } <= download_labels


def _write_workbook(path, sheets):
    """sheets: {sheet_name: (animal_ids, sexes, values)}, same shape as
    test_io_excel.py's _write_sheet helper (kept local since importing
    across test modules is not this suite's convention)."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, (animal_ids, sexes, values) in sheets.items():
        ws = wb.create_sheet(name)
        ws.append(["Animal ID", *animal_ids])
        ws.append(["Sex", *sexes])
        for i, row in enumerate(values, start=1):
            ws.append([i, *row])
    wb.save(path)


def test_unpaired_animal_ids_across_conditions_shows_true_count_and_warning(tmp_path):
    """IMPORTANT fix-wave item 4. Exact review repro: region "R1"'s two
    conditions use entirely disjoint animal IDs (pre={101,102,103},
    TMX={104,105,106} -- numeric, since io_excel.load_workbook_long now
    rejects a non-numeric animal-ID column header as "not a plausible
    animal ID"; letters stood in for these before that check existed).
    `long_df.groupby("region")["animal"].nunique()` -- the old computation
    -- counts 6 distinct animals and reports "6 paired animal(s) per
    region", which is wrong: none of them were measured under both
    conditions. This confirms the caption now shows the true count (0) and
    a plain warning names the mismatch, using a real uploaded workbook
    (not a call directly into io_excel), so the fix is confirmed wired
    into the UI, not just the underlying function."""
    values = [[10.0 + i, 11.0 + i, 12.0 + i] for i in range(20)]
    path = tmp_path / "unpaired.xlsx"
    _write_workbook(
        path,
        {
            "R1 pre": ([101, 102, 103], ["M", "M", "F"], values),
            "R1 TMX": ([104, 105, 106], ["M", "M", "F"], values),
        },
    )
    at = _fresh_app()
    with open(path, "rb") as f:
        content = f.read()
    _excel_uploader(at).set_value((
        "unpaired.xlsx", content,
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ))
    _number_input(at, "Sampling rate (Hz)").set_value(2.0)
    _radio(at, ["Area", "Diameter"]).set_value("Area")
    at.run()
    assert not at.exception

    captions = [c.value for c in at.caption]
    assert any("0 paired animal(s) per region" in c for c in captions)
    warnings = [w.value for w in at.warning]
    assert any(
        "R1" in w and "do not share the same animal IDs" in w for w in warnings
    )


def test_paired_animal_ids_across_conditions_shows_no_warning(tmp_path):
    """Companion negative case: the same shape, but both conditions
    genuinely share their animal IDs -- must report the true count (3)
    and raise no pairing warning, so the fix does not cry wolf on the
    common, correct case. Numeric IDs for the same reason as the sibling
    test above (io_excel.load_workbook_long requires a plausible, numeric
    animal-ID column header)."""
    values = [[10.0 + i, 11.0 + i, 12.0 + i] for i in range(20)]
    path = tmp_path / "paired.xlsx"
    _write_workbook(
        path,
        {
            "R1 pre": ([101, 102, 103], ["M", "M", "F"], values),
            "R1 TMX": ([101, 102, 103], ["M", "M", "F"], values),
        },
    )
    at = _fresh_app()
    with open(path, "rb") as f:
        content = f.read()
    _excel_uploader(at).set_value((
        "paired.xlsx", content,
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ))
    _number_input(at, "Sampling rate (Hz)").set_value(2.0)
    _radio(at, ["Area", "Diameter"]).set_value("Area")
    at.run()
    assert not at.exception

    captions = [c.value for c in at.caption]
    assert any("3 paired animal(s) per region" in c for c in captions)
    warnings = [w.value for w in at.warning]
    assert not any("do not share the same animal IDs" in w for w in warnings)


# --- fix-wave item 3: signal_type/fs no-default + run_pipeline wiring ---

@requires_sample_file
def test_pipeline_does_not_run_until_signal_type_is_chosen():
    """Mutation guard: `signal_type = signal_type_label.lower() if
    signal_type_label else None` (app.py) must never fall back to a
    truthy default such as "area" -- that would run the pipeline (with
    area's square root silently applied) before the user has ever chosen
    Area vs Diameter. test_app_signal_type_radio_has_no_default only
    checks the RADIO WIDGET's own displayed value, which a mutation to
    the `else` branch (e.g. `else "area"`) leaves completely untouched --
    this drives the one path that mutation would actually break: file
    uploaded, fs set, radio deliberately left untouched."""
    at = _fresh_app()
    with open(SAMPLE_PATH, "rb") as f:
        content = f.read()
    _excel_uploader(at).set_value((
        "Vasomotion analysis compilation_Bal K.xlsx", content,
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ))
    _number_input(at, "Sampling rate (Hz)").set_value(FS)
    at.run()
    assert not at.exception
    assert _radio(at, ["Area", "Diameter"]).value is None
    info_texts = [i.value for i in at.info]
    assert any("Choose whether the signal is Area or Diameter" in t for t in info_texts)
    assert "Group-average spectra" not in [h.value for h in at.header]


@requires_sample_file
def test_pipeline_does_not_run_until_sampling_rate_is_entered():
    """Mutation guard: `st.session_state.setdefault("fs_input", 0.0)`
    must never seed a plausible-looking default such as 2.0 -- fs is one
    of the spec's two flagship no-default inputs, and 2.0 happens to be
    this very sample file's own fixed fs (FS above), so a "2.0 default"
    mutation would be invisible to every other test in this file (all of
    which explicitly set fs=2.0 anyway before doing anything else). This
    drives the one path that would expose it: file uploaded, signal type
    chosen, fs deliberately never touched."""
    at = _fresh_app()
    with open(SAMPLE_PATH, "rb") as f:
        content = f.read()
    _excel_uploader(at).set_value((
        "Vasomotion analysis compilation_Bal K.xlsx", content,
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ))
    _radio(at, ["Area", "Diameter"]).set_value("Area")
    at.run()
    assert not at.exception
    assert _number_input(at, "Sampling rate (Hz)").value == 0.0
    info_texts = [i.value for i in at.info]
    assert any("Enter a sampling rate to begin" in t for t in info_texts)
    assert "Group-average spectra" not in [h.value for h in at.header]


@requires_sample_file
def test_signal_type_passed_to_run_pipeline_matches_the_radio(monkeypatch):
    """Item 3's central guard, and the most important of the three: the
    value app.py hands `bundle.run_pipeline` as `signal_type` must be
    exactly what the Signal-is radio says, never a hardcoded literal
    (e.g. `signal_type="area"` regardless of the widget) -- the exact
    mutation the reviewer probed for. That mutation is invisible to every
    other assertion in this suite: the radio widget still shows whatever
    the user picked, parameters.json's "signal_type" field still reads
    from the same (unmutated) local variable everything else here checks,
    and run_pipeline accepts "area" just as happily as "diameter" and
    still returns full results -- so a test that only checks "did it
    crash" or "what does the widget/JSON say" cannot catch it. This wraps
    the real run_pipeline (not a stub -- results must still be genuine)
    to record exactly what app.py called it with, run once per radio
    option so a hardcoded default in either direction is caught."""
    calls = []
    real_run_pipeline = bundle.run_pipeline

    def spy(*args, **kwargs):
        calls.append(kwargs.get("signal_type"))
        return real_run_pipeline(*args, **kwargs)

    monkeypatch.setattr(bundle, "run_pipeline", spy)

    with open(SAMPLE_PATH, "rb") as f:
        content = f.read()
    for label in ("Area", "Diameter"):
        at = _fresh_app()
        _excel_uploader(at).set_value((
            "Vasomotion analysis compilation_Bal K.xlsx", content,
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        ))
        _number_input(at, "Sampling rate (Hz)").set_value(FS)
        _radio(at, ["Area", "Diameter"]).set_value(label)
        at.run()
        assert not at.exception
        assert "Group-average spectra" in [h.value for h in at.header]

    assert calls == ["area", "diameter"]


# --- fix-wave item 2: high-pass cutoff bounded against Nyquist ----------

@requires_sample_file
def test_highpass_cutoff_at_nyquist_is_not_applied_or_recorded():
    """Item 2 (fix wave). preprocess.highpass_filter silently passes a
    trace through unchanged for any cutoff at or above the Nyquist
    frequency (fs/2) -- exact repro from the brief: at fs=2 Hz (Nyquist=
    1.0 Hz), a too-high cutoff used to leave the checkbox ticked, the
    field showing that number, and parameters.json recording it as though
    applied, while nothing was actually filtered. 1.0 Hz -- the Nyquist
    frequency itself, not something further above it like 1.5 -- is used
    here because the widget's own max_value=nyquist bound (this fix's
    first layer) makes Streamlit silently fall back to min_value for
    anything entered above 1.0 (confirmed empirically against streamlit
    1.63.0's number_input: an out-of-range stale value degrades to
    min_value, not a clamp to max_value or a crash), so app.py's own code
    would never even see a raw 1.5 to warn about. 1.0 is exactly the
    boundary the *second* layer (validate_highpass_cutoff's `>=` check,
    mirroring highpass_filter's own no-op condition) exists to catch,
    since bounding the widget cannot exclude its own boundary value.
    Confirms a clear warning names the Nyquist frequency and that the
    value which would actually reach run_pipeline and parameters.json
    (_test_highpass_cutoff -- st.download_button exposes no more to
    AppTest than whether it was clicked) is None, not 1.0: an inert
    filter must never be recorded as applied."""
    at = _upload_sample_and_configure(_fresh_app())
    _number_input(at, "High-pass cutoff (Hz)").set_value(1.0)
    at.run()
    assert not at.exception

    # The widget itself accepted 1.0 (proof max_value really is fs/2 = 1.0,
    # not something stricter that would have rejected it before app.py's
    # own validation ever ran).
    assert _number_input(at, "High-pass cutoff (Hz)").value == 1.0
    assert _checkbox(at, "High-pass filter").value is True

    warnings = [w.value for w in at.warning]
    assert any("Nyquist" in w and "1.0" in w for w in warnings)
    assert at.session_state["_test_highpass_cutoff"] is None


@requires_sample_file
def test_highpass_cutoff_well_below_nyquist_is_applied_and_recorded():
    """Negative companion: the shipped default cutoff (0.005 Hz) at
    fs=2 Hz is nowhere near the 1.0 Hz Nyquist frequency, so item 2's fix
    must not cry wolf on the common, correct case -- no warning, and the
    real value (not None) reaches _test_highpass_cutoff."""
    at = _upload_sample_and_configure(_fresh_app())
    assert not at.exception
    assert at.session_state["_test_highpass_cutoff"] == 0.005
    warnings = [w.value for w in at.warning]
    assert not any("Nyquist" in w for w in warnings)


# --- statistics, duplicate traces, and per-animal xscale ----------------

def test_statistics_section_renders_paired_and_region_tables():
    """The statistics tables reach the UI, with the direction of the
    comparison stated ('pre' as baseline, not the alphabetically-first
    'TMX')."""
    at = _upload_sample_and_configure(_fresh_app())
    assert not at.exception

    headers = [h.value for h in at.header]
    assert "Statistics" in headers

    subheaders = [s.value for s in at.subheader]
    assert any("pre vs TMX" in s for s in subheaders), subheaders
    assert any("Between regions" in s for s in subheaders), subheaders

    # Both statistics dataframes are rendered, alongside the band summary,
    # full PSD and per-trace tables.
    rendered = [df.value for df in at.dataframe]
    paired = [
        d for d in rendered
        if "wilcoxon_p" in getattr(d, "columns", []) and "region" in getattr(d, "columns", [])
    ]
    assert len(paired) == 1, "expected exactly one paired-statistics table"
    table = paired[0]
    # 2 regions x 2 default bands.
    assert len(table) == 4
    assert set(table["condition_baseline"]) == {"pre"}
    assert set(table["condition_treatment"]) == {"TMX"}

    region_tables = [d for d in rendered if "hedges_g" in getattr(d, "columns", [])]
    assert len(region_tables) == 1
    assert len(region_tables[0]) == 2  # one row per band


def test_statistics_download_buttons_are_offered():
    at = _upload_sample_and_configure(_fresh_app())
    labels = [b.label for b in at.download_button]
    assert any("paired statistics" in label for label in labels), labels
    assert any("between-region statistics" in label for label in labels), labels


def test_real_workbook_duplicate_column_is_surfaced_as_an_error():
    """The animal 461672 defect must be visible in the UI, not merely
    detectable by a function: its SpVB TMX column duplicates its SpVB pre
    column, which every other check in the app reports as clean."""
    at = _upload_sample_and_configure(_fresh_app())
    assert not at.exception

    errors = [e.value for e in at.error]
    assert any("Identical traces detected" in e for e in errors), errors

    warnings = [w.value for w in at.warning]
    assert any("461672" in w and "identical values in all 1173 samples" in w for w in warnings)

    # And the consequence is stated where the numbers are, not only at the top.
    assert any("n_zero_diff" in e for e in errors), errors


def test_zero_difference_animal_is_counted_in_the_paired_table():
    at = _upload_sample_and_configure(_fresh_app())
    rendered = [df.value for df in at.dataframe]
    paired = _one(
        [d for d in rendered if "n_zero_diff" in getattr(d, "columns", [])],
        "the paired-statistics table",
    )
    spvb = paired[paired["region"] == "SpVB"]
    smfp = paired[paired["region"] == "SmFP"]
    # 461672 is in SpVB only, so exactly that region shows the zero pair.
    assert (spvb["n_zero_diff"] == 1).all()
    assert (smfp["n_zero_diff"] == 0).all()
    # Dropping the zero pair raises the achievable Wilcoxon floor.
    assert (spvb["wilcoxon_min_p"].to_numpy() > smfp["wilcoxon_min_p"].to_numpy()).all()


def test_wilcoxon_floor_is_explained_at_small_n():
    at = _upload_sample_and_configure(_fresh_app())
    infos = [i.value for i in at.info]
    assert any(
        "cannot return a p-value below" in i and "statement about n" in i for i in infos
    ), infos


def test_frequency_axis_radio_reaches_the_per_animal_figures():
    """Companion to the group-average regression guard: the per-animal small
    multiples previously stayed locked to a log frequency axis while the
    group figure above them switched to linear."""
    at = _upload_sample_and_configure(_fresh_app())
    figures_by_kind = at.session_state["_test_figures"]
    for region_fig in figures_by_kind["per_animal"].values():
        visible = [ax for ax in region_fig.axes if ax.get_visible()]
        assert visible
        assert all(ax.get_xscale() == "log" for ax in visible)

    _radio(at, ["log", "linear"]).set_value("linear")
    at.run()
    assert not at.exception
    figures_by_kind = at.session_state["_test_figures"]
    for region_fig in figures_by_kind["per_animal"].values():
        visible = [ax for ax in region_fig.axes if ax.get_visible()]
        assert visible
        assert all(ax.get_xscale() == "linear" for ax in visible), "per-animal axis ignored xscale"
        assert all(ax.get_yscale() == "log" for ax in visible), "y axis must stay log"


def test_non_normalized_run_warns_that_power_carries_vessel_size():
    at = _fresh_app()
    at = _upload_sample_and_configure(at)
    _selectbox(at, "Normalise").set_value("none")
    at.run()
    assert not at.exception
    warnings = [w.value for w in at.warning]
    assert any(
        "Normalisation is off" in w and "baseline vessel diameter" in w for w in warnings
    ), warnings


def test_single_condition_workbook_reports_no_paired_comparison(tmp_path):
    """A one-condition workbook is valid input for the spectra and band
    tables; it just has no paired contrast. That must read as information,
    not as an error or a crash."""
    values = [[10.0 + i, 11.0 + i, 12.0 + i] for i in range(300)]
    path = tmp_path / "single_condition.xlsx"
    _write_workbook(path, {"R1 pre": ([101, 102, 103], ["M", "M", "F"], values)})
    at = _fresh_app()
    with open(path, "rb") as f:
        content = f.read()
    _excel_uploader(at).set_value((
        "single_condition.xlsx", content,
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ))
    _number_input(at, "Sampling rate (Hz)").set_value(2.0)
    _radio(at, ["Area", "Diameter"]).set_value("Area")
    at.run()
    assert not at.exception
    infos = [i.value for i in at.info]
    assert any("No paired comparison available" in i for i in infos), infos


def test_app_passes_both_statistics_tables_into_the_results_zip():
    """AppTest's DownloadButton does not expose the file bytes, so the zip's
    contents cannot be read back through the harness (its member list is
    locked in test_bundle.py instead). What is checked here is the wiring:
    app.py must actually forward both statistics tables to
    build_results_zip, not compute them for on-screen display only."""
    source = inspect.getsource(app_module)
    call_start = source.index("bundle_module.build_results_zip(")
    call = source[call_start:source.index(")", source.index("parameters", call_start))]
    assert "stats_paired=stats_paired" in call, call
    assert "stats_region=stats_region" in call, call
    assert "trace_stats=trace_stats" in call, call


def test_dockerfile_copies_every_top_level_module():
    """Guards a deploy-only failure mode with no local symptom.

    The Dockerfile previously listed each module by name. Adding stats.py and
    importing it from app.py left the whole test suite green while the built
    container died at startup with ModuleNotFoundError, because the COPY list
    is the only place that knowledge lived and nothing verified it. This test
    fails in CI instead of in production: every top-level module in
    vasomotion_app/ must be reachable through the Dockerfile's COPY
    directives, whether they are spelled out or globbed."""
    import fnmatch

    app_dir = os.path.dirname(os.path.abspath(app_module.__file__))
    dockerfile = os.path.join(app_dir, "Dockerfile")
    with open(dockerfile) as f:
        copy_lines = [
            line for line in f if line.strip().startswith("COPY")
            and "/app/vasomotion_app/" in line
        ]
    assert copy_lines, "no COPY into /app/vasomotion_app/ found in the Dockerfile"

    patterns = []
    for line in copy_lines:
        # "COPY a.py b.py /app/vasomotion_app/" -> the sources are everything
        # between the COPY keyword and the final destination argument.
        parts = line.split()[1:-1]
        patterns.extend(parts)

    modules = sorted(
        name for name in os.listdir(app_dir)
        if name.endswith(".py") and os.path.isfile(os.path.join(app_dir, name))
    )
    assert "stats.py" in modules, "precondition: stats.py exists in the package"

    missing = [
        module for module in modules
        if not any(fnmatch.fnmatch(module, pattern) for pattern in patterns)
    ]
    assert not missing, (
        f"module(s) {missing} are never copied into the container image, so the deployed "
        f"app will fail to import them. Dockerfile COPY sources: {patterns}"
    )
