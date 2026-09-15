"""Tests for vasomotion_app.io_excel."""
from __future__ import annotations

import math
from pathlib import Path

import openpyxl
import pandas as pd
import pytest

from vasomotion_app.io_excel import (
    NonNumericDataError,
    check_animal_pairing,
    detect_sheet_structure,
    load_workbook_long,
    parse_sheet_name,
)

# The real sample workbook this tool is built for. Tests that read it are
# skipped (not failed) in a checkout that doesn't have the data directory.
REAL_WORKBOOK = (
    Path(__file__).resolve().parents[2]
    / "data"
    / "Tim"
    / "Vasomotion analysis compilation_Bal K.xlsx"
)


def _write_sheet(wb, name, animal_ids, sexes, values):
    ws = wb.create_sheet(name)
    ws.append(["Animal ID", *animal_ids])
    ws.append(["Sex", *sexes])
    for i, row in enumerate(values, start=1):
        ws.append([i, *row])


@pytest.fixture
def simple_workbook(tmp_path):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    _write_sheet(
        wb, "RegionA condA",
        animal_ids=[101, 102],
        sexes=["M", "F"],
        values=[[10, 20], [11, 21], [12, 22]],
    )
    _write_sheet(
        wb, "RegionA condB",
        animal_ids=[101, 102],
        sexes=["M", "F"],
        values=[[30, 40], [31, 41], [32, 42]],
    )
    path = tmp_path / "simple.xlsx"
    wb.save(path)
    return str(path)


def test_parse_sheet_name_valid():
    assert parse_sheet_name("SpVB pre") == ("SpVB", "pre")
    assert parse_sheet_name("SmFP TMX") == ("SmFP", "TMX")


def test_parse_sheet_name_invalid():
    assert parse_sheet_name("NoSpaceHere") is None
    assert parse_sheet_name("Too Many Words Here") is None


def test_detect_sheet_structure_all_parsed():
    result = detect_sheet_structure(["SpVB pre", "SpVB TMX", "SmFP pre", "SmFP TMX"])
    assert result["all_parsed"] is True
    assert result["regions"] == ["SmFP", "SpVB"]
    assert result["conditions"] == ["TMX", "pre"]
    assert result["parsed"]["SpVB pre"] == ("SpVB", "pre")


def test_detect_sheet_structure_unparseable():
    result = detect_sheet_structure(["SpVB pre", "WeirdSheetName"])
    assert result["all_parsed"] is False
    assert result["parsed"]["WeirdSheetName"] is None
    assert result["parsed"]["SpVB pre"] == ("SpVB", "pre")


def test_load_workbook_long_shape_and_values(simple_workbook):
    df = load_workbook_long(simple_workbook)
    assert list(df.columns) == ["animal", "sex", "region", "condition", "sample", "value"]
    assert len(df) == 2 * 2 * 3  # 2 sheets x 2 animals x 3 samples

    row = df[(df.animal == 101) & (df.condition == "condA") & (df["sample"] == 1)]
    assert len(row) == 1
    assert row.iloc[0]["value"] == 10.0
    assert row.iloc[0]["sex"] == "M"
    assert row.iloc[0]["region"] == "RegionA"

    row2 = df[(df.animal == 102) & (df.condition == "condB") & (df["sample"] == 3)]
    assert len(row2) == 1
    assert row2.iloc[0]["value"] == 42.0


def test_load_workbook_long_unparseable_sheet_raises(tmp_path):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    _write_sheet(wb, "NotAValidName", [101], ["M"], [[5], [6]])
    path = tmp_path / "bad.xlsx"
    wb.save(path)
    with pytest.raises(ValueError, match="NotAValidName"):
        load_workbook_long(str(path))


def test_load_workbook_long_sheet_regions_override(tmp_path):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    _write_sheet(wb, "NotAValidName", [101], ["M"], [[5], [6]])
    path = tmp_path / "bad.xlsx"
    wb.save(path)
    df = load_workbook_long(str(path), sheet_regions={"NotAValidName": ("RegionX", "condY")})
    assert len(df) == 2
    assert df.iloc[0]["region"] == "RegionX"
    assert df.iloc[0]["condition"] == "condY"


def test_load_workbook_long_non_numeric_cell_raises(tmp_path):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    _write_sheet(wb, "RegionA condA", [101, 102], ["M", "F"], [[10, "oops"], [11, 21]])
    path = tmp_path / "nonnumeric.xlsx"
    wb.save(path)
    with pytest.raises(NonNumericDataError) as excinfo:
        load_workbook_long(str(path))
    assert len(excinfo.value.violations) == 1
    v = excinfo.value.violations[0]
    assert v["sheet"] == "RegionA condA"
    assert v["animal_id"] == 102
    assert v["row"] == 1
    assert v["value"] == "oops"


def test_load_workbook_long_blank_cell_raises(tmp_path):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    _write_sheet(wb, "RegionA condA", [101, 102], ["M", "F"], [[10, None], [11, 21]])
    path = tmp_path / "blank.xlsx"
    wb.save(path)
    with pytest.raises(NonNumericDataError) as excinfo:
        load_workbook_long(str(path))
    assert len(excinfo.value.violations) == 1
    v = excinfo.value.violations[0]
    assert v["sheet"] == "RegionA condA"
    assert v["animal_id"] == 102
    assert v["row"] == 1
    assert math.isnan(v["value"])


# ---------------------------------------------------------------------------
# ITEM 1: trailing 'Mean'/'SD' summary rows and columns must not corrupt data
# ---------------------------------------------------------------------------


def test_load_workbook_long_trailing_summary_row_is_reported_not_crashed(tmp_path):
    """The row case. A trailing 'Mean'/'SD' row has a non-numeric label in
    column A but numeric values underneath (they are computed statistics),
    so this cannot be caught by the per-value numeric check alone -- before
    the fix, `int("Mean")` raised an uncaught ValueError that app.py's
    `except NonNumericDataError` does not catch."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("RegionA condA")
    ws.append(["Animal ID", 101, 102])
    ws.append(["Sex", "M", "F"])
    ws.append([1, 10, 20])
    ws.append([2, 11, 21])
    ws.append([3, 12, 22])
    ws.append(["Mean", 11.0, 21.0])
    ws.append(["SD", 1.0, 1.0])
    path = tmp_path / "trailing_row.xlsx"
    wb.save(path)

    with pytest.raises(NonNumericDataError) as excinfo:
        load_workbook_long(str(path))
    # Every column of both summary rows is reported (2 animals x 2 rows) --
    # never crashes, and never silently folded in as real sample data.
    assert len(excinfo.value.violations) == 4
    assert {v["row"] for v in excinfo.value.violations} == {"Mean", "SD"}
    assert all("row label" in v["reason"] for v in excinfo.value.violations)


def test_load_workbook_long_trailing_summary_column_excluded_not_averaged_in(tmp_path):
    """The column case (the serious one). A trailing 'Mean'/'SD' column has
    a non-numeric header but numeric values underneath (computed
    statistics), so it must never be counted as a real animal or folded
    into a group average -- before the fix this passed silently and
    'Mean'/'SD' showed up as extra 'animals'."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("RegionA condA")
    ws.append(["Animal ID", 101, 102, "Mean", "SD"])
    ws.append(["Sex", "M", "F", None, None])
    ws.append([1, 10, 20, 15.0, 5.0])
    ws.append([2, 11, 21, 16.0, 5.0])
    ws.append([3, 12, 22, 17.0, 5.0])
    path = tmp_path / "trailing_col.xlsx"
    wb.save(path)

    with pytest.raises(NonNumericDataError) as excinfo:
        load_workbook_long(str(path))
    # Every sample row of both summary columns is reported (2 columns x 3
    # samples). Critically, 'Mean'/'SD' never reach a returned dataframe as
    # an "animal" -- the load raises instead of returning contaminated data.
    assert len(excinfo.value.violations) == 6
    assert {v["animal_id"] for v in excinfo.value.violations} == {"Mean", "SD"}
    assert all("plausible animal ID" in v["reason"] for v in excinfo.value.violations)


def test_load_workbook_long_large_integer_animal_id_is_accepted(tmp_path):
    """Regression guard for the plausibility check above: real animal IDs
    are 6-digit integers like 371146 (per the real workbook), not small
    sequential numbers -- the fix must not start rejecting those."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    _write_sheet(
        wb, "RegionA condA",
        animal_ids=[371146, 410328],
        sexes=["M", "F"],
        values=[[10, 20], [11, 21], [12, 22]],
    )
    path = tmp_path / "real_ids.xlsx"
    wb.save(path)
    df = load_workbook_long(str(path))
    assert set(df["animal"].unique()) == {371146, 410328}


# ---------------------------------------------------------------------------
# ITEM 2: NonNumericDataError's message must be capped, however many
# violations there are (fixed entirely inside the exception class, since
# app.py -- which renders str(exc) verbatim -- belongs to another agent).
# ---------------------------------------------------------------------------


def _many_violations(n):
    return [
        {"sheet": "S", "animal_id": float("nan"), "row": i, "value": float("nan")}
        for i in range(n)
    ]


def test_non_numeric_data_error_message_is_capped_with_many_violations():
    # Real reported bug: one blank animal column across two sheets produced
    # a 119,816-character error box. Reproduce the shape (many violations)
    # directly against the exception class, which is where the fix lives.
    violations = _many_violations(5000)
    exc = NonNumericDataError(violations)
    message = str(exc)
    assert len(message) < 5000
    assert "5000" in message  # the true total is still stated
    assert "more" in message.lower()  # says how many were omitted
    # Nothing is lost programmatically even though the message is capped.
    assert len(exc.violations) == 5000


def test_non_numeric_data_error_message_not_flagged_omitted_when_few_violations():
    violations = _many_violations(2)
    exc = NonNumericDataError(violations)
    message = str(exc)
    assert "more" not in message.lower()
    assert len(exc.violations) == 2


def test_load_workbook_long_many_violations_message_is_capped_end_to_end(tmp_path):
    """The exact real-world shape: a blank animal-ID column turns every
    sample row in that column into its own violation. Recreated through the
    real loader (not just the exception class in isolation) to prove the
    cap holds on the actual path app.py's `st.error(str(exc))` renders.

    The column carries real numeric values below the blank header (an
    all-blank column would be trimmed away entirely by openpyxl/pandas
    before `load_workbook_long` ever sees it) -- exactly the "computed
    statistic under a bad header" shape from the ITEM 1 column case, just
    scaled up to many rows instead of a handful."""
    n_samples = 500
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("RegionA condA")
    ws.append(["Animal ID", 101, ""])  # second animal column has a blank header
    ws.append(["Sex", "M", ""])
    for i in range(1, n_samples + 1):
        ws.append([i, 10 + i, 20 + i])
    path = tmp_path / "many_violations.xlsx"
    wb.save(path)

    with pytest.raises(NonNumericDataError) as excinfo:
        load_workbook_long(str(path))
    assert len(excinfo.value.violations) == n_samples
    message = str(excinfo.value)
    assert len(message) < 5000
    assert "more" in message.lower()


# ---------------------------------------------------------------------------
# check_animal_pairing
# ---------------------------------------------------------------------------


def _long_df(rows):
    """rows: list of (animal, region, condition) -- 'sex'/'sample'/'value'
    are irrelevant to pairing and filled with placeholders."""
    return pd.DataFrame(
        [
            {"animal": a, "sex": "M", "region": r, "condition": c, "sample": 1, "value": 0.0}
            for a, r, c in rows
        ]
    )


def test_check_animal_pairing_genuinely_unpaired_fixture():
    """IMPORTANT fix-wave item 4. The exact scenario from the review:
    pre={A,B,C}, TMX={D,E,F} for the same region -- 6 distinct animals by
    count, 0 actually measured under both conditions. Before this fix,
    `groupby("region")["animal"].nunique()` reported 6 "paired" here,
    which is the one number this check exists to catch as wrong."""
    df = _long_df(
        [(a, "R1", "pre") for a in ["A", "B", "C"]]
        + [(a, "R1", "TMX") for a in ["D", "E", "F"]]
    )
    result = check_animal_pairing(df)
    assert result["n_paired_per_region"] == {"R1": 0}
    assert len(result["warnings"]) == 1
    assert "R1" in result["warnings"][0]
    assert "do not share the same animal IDs" in result["warnings"][0]


def test_check_animal_pairing_fully_paired_no_warning():
    df = _long_df(
        [(a, "R1", "pre") for a in ["A", "B", "C"]]
        + [(a, "R1", "TMX") for a in ["A", "B", "C"]]
    )
    result = check_animal_pairing(df)
    assert result["n_paired_per_region"] == {"R1": 3}
    assert result["warnings"] == []


def test_check_animal_pairing_partial_overlap_counts_only_the_intersection():
    # pre={A,B,C}, TMX={B,C,D}: 2 animals (B, C) genuinely common to both,
    # not the 4 unique animals nunique() would count, and not 0 either --
    # a distinct case from "fully paired" and "fully unpaired".
    df = _long_df(
        [(a, "R1", "pre") for a in ["A", "B", "C"]]
        + [(a, "R1", "TMX") for a in ["B", "C", "D"]]
    )
    result = check_animal_pairing(df)
    assert result["n_paired_per_region"] == {"R1": 2}
    assert len(result["warnings"]) == 1


def test_check_animal_pairing_single_condition_region_is_trivially_paired():
    # A region with only one condition has nothing to compare against --
    # every animal present counts as paired, and no warning is meaningful.
    df = _long_df([(a, "R1", "pre") for a in ["A", "B", "C"]])
    result = check_animal_pairing(df)
    assert result["n_paired_per_region"] == {"R1": 3}
    assert result["warnings"] == []


def test_check_animal_pairing_reports_per_region_independently():
    # R1 is unpaired, R2 is fully paired -- one region's mismatch must not
    # taint or hide the other's correct count.
    df = _long_df(
        [(a, "R1", "pre") for a in ["A", "B"]]
        + [(a, "R1", "TMX") for a in ["C", "D"]]
        + [(a, "R2", "pre") for a in ["X", "Y"]]
        + [(a, "R2", "TMX") for a in ["X", "Y"]]
    )
    result = check_animal_pairing(df)
    assert result["n_paired_per_region"] == {"R1": 0, "R2": 2}
    assert len(result["warnings"]) == 1
    assert "R1" in result["warnings"][0]
    assert "R2" not in result["warnings"][0]


def test_check_animal_pairing_from_real_workbook_via_load_workbook_long(simple_workbook):
    # End-to-end sanity check through the real loader, not just a hand-built
    # DataFrame: simple_workbook's two sheets share animal_ids=[101, 102].
    df = load_workbook_long(simple_workbook)
    result = check_animal_pairing(df)
    assert result["n_paired_per_region"] == {"RegionA": 2}
    assert result["warnings"] == []


@pytest.mark.skipif(
    not REAL_WORKBOOK.exists(),
    reason="sample workbook not present in this checkout",
)
def test_real_sample_file_sheet_names_and_shape():
    wb = openpyxl.load_workbook(str(REAL_WORKBOOK), read_only=True)
    try:
        sheet_names = wb.sheetnames
    finally:
        wb.close()
    assert sheet_names == ["SpVB pre", "SpVB TMX", "SmFP pre", "SmFP TMX"]

    df = load_workbook_long(str(REAL_WORKBOOK))
    assert list(df.columns) == ["animal", "sex", "region", "condition", "sample", "value"]
    assert df.shape == (28152, 6)
